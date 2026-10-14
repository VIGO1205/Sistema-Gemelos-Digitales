"""Trazabilidad y auditoria de los archivos finales, sin tocar produccion."""
from pathlib import Path
import json,csv,hashlib,platform,shutil,re
from datetime import datetime,timezone
import pandas as pd
import pymupdf
from docx import Document
ROOT=Path(__file__).resolve().parents[2];OUT=ROOT/'entrega/articulo_RIAI';E=OUT/'evidencia';T=OUT/'tablas'
run=json.loads((E/'ejecucion.json').read_text(encoding='utf-8'))
launch=run['command']
manifest={'actual_command':launch,'internal_argv':'python experimentos/articulo_riai/evaluar_cinco.py','note':'Revisión con cinco modelos principales y dos referencias. Contenedor aislado; artefactos de producción en solo lectura.','physical_cpu':'AMD Ryzen 7 5700U with Radeon Graphics','physical_cores':8,'logical_processors':16,'retrieved_via':'Get-CimInstance Win32_Processor','utc_recorded':run['utc'],'runtime_seconds':run['duration_seconds'],'versions':run['versions']}
(E/'lanzador_real.json').write_text(json.dumps(manifest,ensure_ascii=False,indent=2),encoding='utf-8')
for name in ['copper_creek_winner.metadata.json']:
 shutil.copy2(ROOT/'motor_ia_m6/model_artifacts'/name,E/name)
scene=json.loads((ROOT/'backend/data/copper_creek_mammoth.json').read_text(encoding='utf-8'))
(E/'escena_metadata.json').write_text(json.dumps(scene['metadata'],ensure_ascii=False,indent=2),encoding='utf-8')
q=json.loads((E/'calidad.json').read_text(encoding='utf-8'))
rows=[]
def tr(label,value,file,seed='42',command=launch):rows.append({'cifra':label,'valor':value,'archivo_de_salida':file,'comando':command,'semilla':seed})
def flatten(x,prefix,file):
 for k,v in x.items():
  if isinstance(v,dict):flatten(v,prefix+'/'+k,file)
  elif isinstance(v,(int,float)):tr(prefix+'/'+k,v,file)
flatten(q,'calidad','evidencia/calidad.json')
for fn in ['medias_semillas.csv','metricas_agrupadas.csv','bootstrap_pareado.csv','cobertura.csv','error_distancia.csv','registro_ajustes.csv','predicciones_negativas.csv','pesos_hibridos.csv']:
 data=pd.read_csv(T/fn)
 for i,r in data.iterrows():
  key=';'.join(f'{k}={r[k]}' for k in ['protocol','model','fold','comparison'] if k in r.index)
  for k,v in r.items():
   if isinstance(v,(int,float)) and pd.notna(v):tr(f'{fn}:fila={i+2};{key};{k}',v,'tablas/'+fn,str(r.get('seed',42)))
cov=pd.read_csv(T/'cobertura.csv')
for name,g in cov.groupby('protocol'):
 tr(name+'/cobertura_pct',100*g.supported_samples.sum()/g.test_samples.sum(),'tablas/cobertura.csv')
 tr(name+'/sin_soporte',int((g.test_samples-g.supported_samples).sum()),'tablas/cobertura.csv')
summary=pd.read_csv(T/'medias_semillas.csv');g=summary[summary.protocol=='sondajes']
for model,d in g.groupby('model'):
 for m in ['rmse','mae','r2']:
  tr(model+'/'+m+'/media_semillas',d[m].mean(),'tablas/medias_semillas.csv','42–46')
  if len(d)>1:tr(model+'/'+m+'/sd_semillas',d[m].std(ddof=1),'tablas/medias_semillas.csv','42–46')
tr('RF_R2_grupo_4',pd.read_csv(T/'metricas_pliegues.csv').query("protocol=='sondajes' and model=='Random Forest' and seed==42 and fold==4").r2.iloc[0],'tablas/metricas_pliegues.csv')
tr('numero_cubos',len(scene['blocks']),'backend/data/copper_creek_mammoth.json','n/a','lectura del artefacto publicado')
tr('niveles',len(scene['pit_levels']),'backend/data/copper_creek_mammoth.json','n/a','lectura del artefacto publicado')
for k in ['bench_height_m','skipped_unsupported_blocks','displayed_holes']:tr(k,scene['metadata'][k],'evidencia/escena_metadata.json','n/a','lectura del artefacto publicado')
params={
'ventana_collares_x_min':547650,'ventana_collares_x_max':549300,'ventana_collares_y_min':3623100,'ventana_collares_y_max':3624750,
'ventana_muestras_x_min':547950,'ventana_muestras_x_max':548990,'ventana_muestras_y_min':3623410,'ventana_muestras_y_max':3624430,'ventana_muestras_z_min':1120,'ventana_muestras_z_max':1350,
'cu_admisible_min':0,'cu_admisible_max':20,'integracion_m':10,'composite_nominal_m':6,'composite_min_m':1,
'origen_modelo_x':548470,'origen_modelo_y':3623920,'origen_modelo_z':1250,'rf_arboles':160,'rf_hoja_min':4,'rf_variables':2,
'xgb_arboles':350,'xgb_profundidad':4,'xgb_tasa':.05,'xgb_submuestreo':.8,'xgb_columnas':.9,'trabajadores':4,'idw_exponente':3,'radio_m':120,'min_sondajes':3,'max_sondajes':8,'idw_distancia_min_m':1,'candidatos_vecinos':160,
'particiones':5,'semilla_publicada':42,'numero_semillas':5,'celda_m':200,'origen_celdas_x':548000,'origen_celdas_y':3623400,'bootstrap_replicas':2000,'percentil_inferior':2.5,'percentil_superior':97.5,
'cubo_dx_m':20,'cubo_dy_m':20,'cubo_dz_m':12,'nivel_max_m':1316,'nivel_min_m':1172,'figura_5_saturacion_cu_pct':1,'decimales_escena':4,'consultas_validas_comprobadas':3,'consultas_exteriores_comprobadas':1}
params.update({'modelos_principales':5,'referencias_comparativas':2,'kriging_max_muestras':500,'kriging_vecinos':64,'variograma_intervalos':20,'stacking_reserva_grupos':.3,'stacking_arboles':150,'stacking_ridge_alpha':1,'nn_reserva_grupos':.15,'nn_epocas':100,'nn_dropout':.2,'nn_tasa':.001,'nn_lote':32,'nn_pesos_fusion_candidatos':21,'kriging_semilla_interna_hibridos':42})
for k,v in params.items():tr(k,v,'ficha_tecnica.md; evidencia/ecuaciones_vs_codigo.md; experimentos/articulo_riai/evaluar_cinco.py; motor_ia_m6/models/; backend/scripts/copper_creek_samples.py; backend/scripts/benchmark_copper_creek.py; backend/scripts/build_copper_creek_reconstruction.py','configuración','inspección del código y archivos de evidencia')
selection=json.loads((E/'seleccion_cinco.json').read_text(encoding='utf-8'))
tr('modelo_seleccionado_estudio',selection['study_winner'],'evidencia/seleccion_cinco.json')
for p in sorted((E/'ajustes_cinco').glob('*.json')):
 rec=json.loads(p.read_text(encoding='utf-8'));flatten(rec['details'],p.stem+'/details','evidencia/ajustes_cinco/'+p.name)
for k,v in run['versions'].items():tr('version_'+k,v,'evidencia/ejecucion.json','n/a')
for k,v in [('cpu_logicos',16),('CPU','AMD Ryzen 7 5700U'),('Python','3.13.16')]:tr(k,v,'evidencia/lanzador_real.json; evidencia/ejecucion.json','n/a')
delta=json.loads((E/'reproduccion_metadata.json').read_text(encoding='utf-8'))
tr('diferencia_xgb_rmse_previo',delta['XGBoost']['rmse'],'evidencia/reproduccion_metadata.json')
with (E/'trazabilidad.csv').open('w',encoding='utf-8-sig',newline='') as f:
 w=csv.DictWriter(f,fieldnames=rows[0].keys());w.writeheader();w.writerows(rows)

pdf=pymupdf.open(OUT/'Articulo_RIAI_Copper_Creek.pdf');doc=Document(OUT/'Articulo_RIAI_Copper_Creek.docx')
fonts=[]
for i in range(len(pdf)):
 for ft in pdf[i].get_fonts():
  base=pdf.extract_font(ft[0]);fonts.append({'font':ft[3],'embedded':bool(base[3])})
final_hash={p:hashlib.sha256((ROOT/p).read_bytes()).hexdigest() for p in run['protected_before']}
content='\n'.join(p.get_text() for p in pdf)
footer_numbers=[]
for p in pdf:
 texts=[b[4].strip() for b in p.get_text('words') if b[1]>p.rect.height-60 and re.fullmatch(r'\d+',b[4].strip())]
 footer_numbers.append(texts)
last_content=[b[3] for b in pdf[-1].get_text('blocks') if 45<b[1]<pdf[-1].rect.height-60]
last_fraction=(max(last_content)-56.7)/(pdf[-1].rect.height-113.4)
audit={'pages':len(pdf),'max_pages':12,'equations_native':len(doc._element.xpath('.//m:oMath')),'tables':len(doc.tables),'inline_figures':len(doc.inline_shapes),'all_pdf_fonts_embedded':all(f['embedded'] for f in fonts),'active_pdf_links':sum(len(p.get_links()) for p in pdf),'template_unchanged':hashlib.sha256((ROOT/'plantilla_revista_RIAI/plantilla.docx').read_bytes()).hexdigest()=='5bce0799592bbe719860bbbb8caa26da0e709823cdfc086b2c5de31a95890dc1','production_unchanged':final_hash==run['protected_before'],'unresolved_markers':bool(re.search(r'\{\{(?:TABLE|EQ|REFERENCES)',content)),'page_numbers':footer_numbers,'last_page_body_height_fraction':last_fraction,'traceability_rows':len(rows)}
comparison_names=[r.cells[0].text for r in doc.tables[2].rows[1:]]
audit.update({'comparison_rows':comparison_names,'five_principal_models_plus_two_references':comparison_names==['OK','RF','XGBoost','Stacking','NN+OK','IDW3 (ref.)','Media (ref.)'],'all_original_model_names_in_body':all(n in '\n'.join(p.text for p in doc.paragraphs) for n in selection['principal_models'])})
assert audit['five_principal_models_plus_two_references'] and audit['all_original_model_names_in_body']
audit['body_headers_present']=all('Paredes-Ilquimiche' in p.get_text(clip=pymupdf.Rect(0,0,p.rect.width,50)) for p in list(pdf)[1:])
audit['final_layout_checks_pass']=len(pdf)<=12 and last_fraction>.5 and audit['body_headers_present'] and audit['all_pdf_fonts_embedded'] and audit['active_pdf_links']==0 and not audit['unresolved_markers']
assert audit['final_layout_checks_pass'] and audit['production_unchanged'] and audit['template_unchanged']
(E/'auditoria_final.json').write_text(json.dumps(audit,ensure_ascii=False,indent=2),encoding='utf-8')
files={str(p.relative_to(OUT)):hashlib.sha256(p.read_bytes()).hexdigest() for p in OUT.rglob('*') if p.is_file() and p.name!='manifest_sha256.json' and 'qa' not in p.parts and 'entorno_local' not in p.parts}
(E/'manifest_sha256.json').write_text(json.dumps(files,ensure_ascii=False,indent=2),encoding='utf-8')
print(json.dumps(audit,ensure_ascii=False,indent=2))
