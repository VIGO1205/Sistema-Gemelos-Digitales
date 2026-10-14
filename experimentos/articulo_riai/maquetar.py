"""Maqueta desde la plantilla retenida; sustituye marcadores por evidencia real."""
from pathlib import Path
import re,json,copy,hashlib,csv
from lxml import etree
import pandas as pd
from docx import Document
from docx.shared import Cm,Pt,RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH,WD_BREAK,WD_TAB_ALIGNMENT
from docx.enum.section import WD_SECTION_START
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT=Path(__file__).resolve().parents[2];OUT=ROOT/'entrega/articulo_RIAI';T=OUT/'tablas'
src=OUT/'articulo.md';raw=src.read_text(encoding='utf-8')
means=pd.read_csv(T/'medias_semillas.csv')
main=means[(means.protocol=='sondajes')&(means.seed==42)]
stress=means[(means.protocol=='celdas_200m')&(means.seed==42)]
tables={
 'TABLE1':('Tabla 1: Inventario de datos y muestra analizada.', ['Elemento','Cantidad'],[
 ['Collares públicos','649'],['Estaciones de trayectoria públicas','9 615'],['Intervalos de ensayo públicos','87 733'],['Registros de litología disponibles','9 490'],['Collares en la selección inicial','390'],['Sondajes con composites seleccionados','164'],['Composites del estudio','2 480'],['Composites con longitud menor que 6 m','223']]),
 'TABLE2':('Tabla 2: Configuraciones explícitas de los estimadores.', ['Modelo','Configuración'],[
 ['OK','Kriging ordinario 3D; variograma esférico automático, 20 intervalos; máximo 500 muestras de ajuste y 64 vecinos.'],
 ['RF','160 árboles; mínimo 4 muestras por hoja; 2 variables por división; 4 trabajadores.'],
 ['XGBoost','350 árboles; profundidad 4; tasa 0.05; submuestreo 0.8; fracción de columnas 0.9; objetivo cuadrático; 4 trabajadores.'],
 ['Stacking','OK, RF y XGBoost; 150 árboles por componente; combinación Ridge con penalización 1; reserva interna de grupos del 30 %.'],
 ['NN+OK','Red 3–128–64–32–1; ReLU, normalización por lote y dropout 0.2; Adam, tasa 0.001; 100 épocas; lote 32; reserva interna del 15 %; fusión con OK.'],
 ['IDW3','Exponente 3; radio 120 m; entre 3 y 8 sondajes; distancia mínima para pesos 1 m.'],
 ['Media','Media de Cu del conjunto de entrenamiento, sin utilizar coordenadas.']]),
}
names=['Ordinary Kriging','Random Forest','XGBoost','Stacking Ensemble','NN + Bayesian Kriging','IDW3','Media']
short={'Ordinary Kriging':'OK','Random Forest':'RF','XGBoost':'XGBoost','Stacking Ensemble':'Stacking','NN + Bayesian Kriging':'NN+OK','IDW3':'IDW3 (ref.)','Media':'Media (ref.)'}
for key,df,cap in [('TABLE3',main,'Tabla 3: Promedio por partición al reservar sondajes completos. RMSE y MAE en puntos porcentuales de Cu; semilla 42.'),('TABLE4',stress,'Tabla 4: Promedio por partición al reservar celdas de 200 m. RMSE y MAE en puntos porcentuales de Cu; semilla 42.')]:
 rows=[]
 for name in names:
  r=df[df.model==name].iloc[0];rows.append([short[name],*[f'{r[k]:.4f}' for k in ['rmse','mae','r2']]])
 tables[key]=(cap,['Modelo','RMSE','MAE','R²'],rows)
for key,(cap,headers,rows) in tables.items():
 with (T/(key.lower()+'.csv')).open('w',encoding='utf-8-sig',newline='') as f:
  w=csv.writer(f);w.writerow(headers);w.writerows(rows)

refs=[
'Anvari, K., Benndorf, J., 2025. Real time mining—A review of developments within the last decade. Mining 5(3), 38. DOI: 10.3390/mining5030038.',
'Ausenco Engineering USA South Inc., 2023. Copper Creek Project NI 43-101 technical report and preliminary economic assessment, Arizona, United States of America. Informe preparado para Faraday Copper Corp., fecha efectiva 3 de mayo de 2023.',
'Breiman, L., 2001. Random forests. Machine Learning 45, 5–32. DOI: 10.1023/A:1010933404324.',
'Chen, T., Guestrin, C., 2016. XGBoost: A scalable tree boosting system. Proceedings of the 22nd ACM SIGKDD International Conference on Knowledge Discovery and Data Mining, 785–794. DOI: 10.1145/2939672.2939785.',
'Faraday Copper, 2026. Copper Creek database: Assays, collars, lithology and surveys, January 2026. Data files for download. faradaycopper.com/projects/copper-creek-project/drilling-data-files-for-download/ (consulta: 7 de octubre de 2026).',
'Kritzinger, W., Karner, M., Traar, G., Henjes, J., Sihn, W., 2018. Digital twin in manufacturing: A categorical literature review and classification. IFAC-PapersOnLine 51(11), 1016–1022. DOI: 10.1016/j.ifacol.2018.08.474.',
'Lindi, O. T., Aladejare, A. E., Ozoji, T. M., Ranta, J.-P., 2024. Uncertainty quantification in mineral resource estimation. Natural Resources Research 33, 2503–2526. DOI: 10.1007/s11053-024-10394-6.',
'Maniteja, M., Samanta, G., Gebretsadik, A., Tsae, N. B., Rai, S. S., Fissha, Y., Okada, N., Kawamura, Y., 2025. Advancing iron ore grade estimation: A comparative study of machine learning and ordinary kriging. Minerals 15(2), 131. DOI: 10.3390/min15020131.',
'Roberts, D. R., Bahn, V., Ciuti, S., Boyce, M. S., Elith, J., Guillera-Arroita, G., Hauenstein, S., Lahoz-Monfort, J. J., Schröder, B., Thuiller, W., Warton, D. I., Wintle, B. A., Hartig, F., Dormann, C. F., 2017. Cross-validation strategies for data with temporal, spatial, hierarchical, or phylogenetic structure. Ecography 40, 913–929. DOI: 10.1111/ecog.02881.',
'Tsae, N. B., Adachi, T., Kawamura, Y., 2023. Application of artificial neural network for the prediction of copper ore grade. Minerals 13(5), 658. DOI: 10.3390/min13050658.',
'Wolpert, D. H., 1992. Stacked generalization. Neural Networks 5(2), 241–259. DOI: 10.1016/S0893-6080(05)80023-1.'
]
(OUT/'referencias.md').write_text('# Referencias verificadas\n\n'+'\n\n'.join(refs),encoding='utf-8')

# Native editable Office Math nodes.
def mr(t):
 r=OxmlElement('m:r');pr=OxmlElement('m:rPr');st=OxmlElement('m:sty');st.set(qn('m:val'),'p' if t in ['RMSE','MAE','R','IDW3'] else 'i');pr.append(st);r.append(pr)
 wr=OxmlElement('w:rPr');sz=OxmlElement('w:sz');sz.set(qn('w:val'),'20');wr.append(sz);r.append(wr)
 x=OxmlElement('m:t');x.text=t;r.append(x);return r
def box(tag,items):
 n=OxmlElement(tag)
 for a in items: n.append(mr(a) if isinstance(a,str) else a)
 return n
def frac(a,b): return box('m:f',[box('m:num',a),box('m:den',b)])
def sub(a,b): return box('m:sSub',[box('m:e',[a]),box('m:sub',[b])])
def sup(a,b): return box('m:sSup',[box('m:e',[a]),box('m:sup',[b])])
def sqrt(a):
 p=OxmlElement('m:radPr');h=OxmlElement('m:degHide');h.set(qn('m:val'),'1');p.append(h)
 return box('m:rad',[p,box('m:deg',[]),box('m:e',a)])
def summ(): return sub('∑','i')
def eq_nodes(n):
 if n==1:return [sub('c̄','j'),' = ',frac([summ(),sub('ℓ','ij'),sub('c','i')],[summ(),sub('ℓ','ij')])]
 if n==2:return [sub('ĉ','IDW3'),' = ',frac([summ(),sub('w','i'),sub('c','i')],[summ(),sub('w','i')])]
 if n==3:return ['RMSE = ',sqrt([frac(['1'],['n']),summ(),sup(box('m:d',[box('m:e',[sub('y','i'),' − ',sub('ŷ','i')])]),'2')])]
 if n==4:return ['MAE = ',frac(['1'],['n']),summ(),'|',sub('y','i'),' − ',sub('ŷ','i'),'|']
 if n==5:return [sup('R','2'),' = 1 − ',frac([summ(),sup(box('m:d',[box('m:e',[sub('y','i'),' − ',sub('ŷ','i')])]),'2')],[summ(),sup(box('m:d',[box('m:e',[sub('y','i'),' − ȳ'])]),'2')])]

template=ROOT/'plantilla_revista_RIAI/plantilla.docx';doc=Document(template)
paras=doc.paragraphs
first_sec=paras[20]._p
body_sec=copy.deepcopy(doc.sections[1]._sectPr)
for node in list(doc._element.body):
 if node is first_sec: break
# Delete instructional content after first continuous break; keep original front matter.
seen=False
for node in list(doc._element.body):
 if seen: doc._element.body.remove(node)
 if node is first_sec: seen=True
doc._element.body.append(body_sec)
for s in doc.sections:
 s.page_width=Cm(21);s.page_height=Cm(29.7);s.left_margin=Cm(1.5);s.right_margin=Cm(1.5);s.top_margin=Cm(2);s.bottom_margin=Cm(2)
 s._sectPr.get_or_add_cols() if hasattr(s._sectPr,'get_or_add_cols') else None
cols=body_sec.find(qn('w:cols'));cols.set(qn('w:num'),'2');cols.set(qn('w:space'),'340')
doc.styles['Normal'].font.name='Times New Roman';doc.styles['Normal'].font.size=Pt(10)
bs=doc.styles['icsm_bodytext'];bs.font.size=Pt(10);bs.paragraph_format.space_after=Pt(3);bs.paragraph_format.line_spacing=1.0;bs.paragraph_format.widow_control=True
for sn in ['Tipo Sección RIAI','Tipo Subsección RIAI','Tipo Titulo RIAI','Epígrafe']:
 s=doc.styles[sn];s.font.color.rgb=RGBColor(0,0,0)
for sn in ['Tipo Sección RIAI','Tipo Subsección RIAI']:
 s=doc.styles[sn];s.paragraph_format.keep_with_next=True;s.paragraph_format.space_before=Pt(8);s.paragraph_format.space_after=Pt(5);s.font.size=Pt(10)
doc.styles['icsm_references'].font.size=Pt(8)
doc.settings.odd_and_even_pages_header_footer=False

title=raw.splitlines()[0][2:];summary=raw.split('## Resumen\n\n')[1].split('\n\nPalabras clave:')[0]
en_title='Digital twin prototype for spatial copper grade estimation'
abstract=raw.split('### Abstract\n\n')[1].split('\n\nKeywords:')[0]
kw='Palabras clave:'+raw.split('Palabras clave:')[1].split('\n')[0]
ekw='Keywords:'+raw.split('Keywords:')[1].split('\n')[0]
paras[0].text=title;paras[1].text='';paras[3].text='Paredes-Ilquimiche, J. I.*, Vigo-Huayán, A. G. G.'
paras[4].text='Universidad Nacional de Trujillo, Facultad de Ingeniería, Escuela Profesional de Ingeniería de Sistemas, Av. Juan Pablo II, 13011 Trujillo, Perú.'
for idx in [1,5,6]:
 el=paras[idx]._p;el.getparent().remove(el)
paras[12].text=summary;paras[13].text=kw;paras[15].text=en_title;paras[17].text=abstract;paras[18].text=ekw
for i in [12,17]:
 paras[i].paragraph_format.line_spacing=1;paras[i].paragraph_format.space_after=Pt(5)
 for r in paras[i].runs:r.font.size=Pt(9)

# Replace editorial sample citation in both drawing and VML fallback.
for tx in doc._element.xpath('.//w:txbxContent'):
 ps=tx.findall(qn('w:p'))
 if ps:
  p=ps[0]
  for c in list(p):
   if c.tag!=qn('w:pPr'):p.remove(c)
  r=OxmlElement('w:r');rp=OxmlElement('w:rPr');z=OxmlElement('w:sz');z.set(qn('w:val'),'16');rp.append(z);color=OxmlElement('w:color');color.set(qn('w:val'),'222222');rp.append(color);r.append(rp);t=OxmlElement('w:t');t.text='Paredes-Ilquimiche, J. I., Vigo-Huayán, A. G. G. Digital twin prototype for spatial copper grade estimation. Manuscrito para evaluación en RIAI.';r.append(t);p.append(r)
  for p2 in ps[1:]:tx.remove(p2)
for part in doc.part.package.parts:
 if str(part.partname).startswith('/word/header') or str(part.partname).startswith('/word/footer'):
  root=part._element
  for p in root.xpath('.//w:p'):
   ts=p.xpath('.//w:t');txt=''.join(t.text or '' for t in ts)
   if 'Apellido primer autor' in txt:replacement='Paredes-Ilquimiche, J. I. y Vigo-Huayán, A. G. G. / Manuscrito para evaluación en RIAI'
   elif 'Industrial 00' in txt:replacement='Revista Iberoamericana de Automática e Informática Industrial · Manuscrito para evaluación'
   elif 'autor1@ceautomatica.es' in txt:replacement='* Autor para correspondencia: jiparedesil@unitru.edu.pe'
   else:continue
   if ts:
    ts[0].text=replacement
    for t in ts[1:]:t.text=''

def para(text,style='icsm_bodytext'):
 p=doc.add_paragraph(text,style)
 if style.startswith('Tipo Secci') or style.startswith('Tipo Subsecci'):
  np=OxmlElement('w:numPr');ni=OxmlElement('w:numId');ni.set(qn('w:val'),'0');np.append(ni);p._p.get_or_add_pPr().append(np)
 return p
def table(key):
 cap,heads,rows=tables[key];p=para(cap,'Tipo Tabla RIAI');p.paragraph_format.keep_with_next=True;p.paragraph_format.space_before=Pt(5)
 tb=doc.add_table(rows=1,cols=len(heads));tb.autofit=False
 widths=[3.0,1.88,1.88,1.88] if len(heads)==4 else ([6.25,2.35] if key=='TABLE1' else [1.75,6.85])
 for j,w in enumerate(widths):tb.columns[j].width=Cm(w)
 for j,h in enumerate(heads):tb.rows[0].cells[j].text=h
 for row in rows:
  cs=tb.add_row().cells
  for j,v in enumerate(row):cs[j].text=str(v)
 for i,row in enumerate(tb.rows):
  pr=row._tr.get_or_add_trPr();ns=OxmlElement('w:cantSplit');pr.append(ns)
  if i==0:pr.append(OxmlElement('w:tblHeader'))
  for j,c in enumerate(row.cells):
   c.width=Cm(widths[j]);tcpr=c._tc.get_or_add_tcPr();border=OxmlElement('w:tcBorders')
   for edge in ['top','bottom','left','right']:
    e=OxmlElement('w:'+edge);e.set(qn('w:val'),'single');e.set(qn('w:sz'),'3');e.set(qn('w:color'),'B0B0B0');border.append(e)
   tcpr.append(border)
   if i==0:
    sh=OxmlElement('w:shd');sh.set(qn('w:fill'),'F1F1F1');tcpr.append(sh)
   for p in c.paragraphs:
    p.paragraph_format.space_after=Pt(3);p.paragraph_format.space_before=Pt(3);p.paragraph_format.line_spacing=1;p.paragraph_format.keep_with_next=(i<len(tb.rows)-1)
    p.alignment=WD_ALIGN_PARAGRAPH.LEFT
    for r in p.runs:r.font.name='Times New Roman';r.font.size=Pt(8)
 para('').paragraph_format.space_after=Pt(0)

body=raw.split('## 1 Introducción')[1];body='## 1 Introducción'+body
inrefs=False
for chunk in body.split('\n\n'):
 chunk=chunk.strip()
 if not chunk:continue
 if chunk.startswith('## '):para(chunk[3:],'Tipo Sección RIAI');continue
 if chunk.startswith('### '):para(chunk[4:],'Tipo Subsección RIAI');continue
 if chunk.startswith('{{TABLE'):table(chunk[2:-2]);continue
 if chunk.startswith('{{EQ'):
  n=int(chunk[4:-2]);p=para('');p.alignment=WD_ALIGN_PARAGRAPH.LEFT;p.paragraph_format.first_line_indent=Cm(0);p.paragraph_format.space_before=Pt(3);p.paragraph_format.space_after=Pt(5)
  p.paragraph_format.tab_stops.add_tab_stop(Cm(4.25),WD_TAB_ALIGNMENT.CENTER);p.paragraph_format.tab_stops.add_tab_stop(Cm(8.55),WD_TAB_ALIGNMENT.RIGHT)
  p.add_run('\t');p._p.append(box('m:oMath',eq_nodes(n)));p.add_run('\t('+str(n)+')');continue
 if chunk.startswith('!['):
  path=re.search(r'\]\((.*?)\)',chunk).group(1);p=para('');p.paragraph_format.first_line_indent=Cm(0);p.alignment=WD_ALIGN_PARAGRAPH.CENTER;p.paragraph_format.keep_with_next=True;p.paragraph_format.space_before=Pt(4);p.add_run().add_picture(str(OUT/path),width=Cm(8.5));continue
 if chunk.startswith('Figura '):
  p=para(chunk,'Tipo Figura RIAI');p.paragraph_format.space_after=Pt(7);p.paragraph_format.line_spacing=1;p.alignment=WD_ALIGN_PARAGRAPH.JUSTIFY
  for r in p.runs:r.font.size=Pt(8)
  continue
 if chunk=='{{REFERENCES}}':
  for ref in refs:
   p=para(ref,'icsm_references');p.paragraph_format.left_indent=Cm(.3);p.paragraph_format.first_line_indent=Cm(-.3);p.paragraph_format.space_after=Pt(3);p.paragraph_format.line_spacing=1;p.alignment=WD_ALIGN_PARAGRAPH.JUSTIFY
  continue
 p=para(chunk)

# A matching continuous section balances the last page without switching
# to one column, which would move the empty section onto another page.
end=doc.add_section(WD_SECTION_START.CONTINUOUS)
end._sectPr.find(qn('w:cols')).set(qn('w:num'),'2')
last=doc.paragraphs[-1];last.paragraph_format.space_before=Pt(0);last.paragraph_format.space_after=Pt(0);last.paragraph_format.line_spacing=Pt(1)
for s in list(doc.sections)[:2]:
 s.header.is_linked_to_previous=False
 h=s.header
 for el in list(h._element):h._element.remove(el)
 p=h.add_paragraph('Paredes-Ilquimiche, J. I. y Vigo-Huayán, A. G. G. / Manuscrito para evaluación en RIAI')
 p.alignment=WD_ALIGN_PARAGRAPH.CENTER;p.paragraph_format.space_after=Pt(0)
 for r in p.runs:r.font.name='Times New Roman';r.font.size=Pt(8);r.font.italic=True
 s.header_distance=Cm(1.0) if s!=doc.sections[0] else Cm(1.27)
 if not s.footer.paragraphs:continue
 p=s.footer.paragraphs[0]
 if p.text or p._p.xpath('.//w:fldSimple'):continue
 p.alignment=WD_ALIGN_PARAGRAPH.RIGHT
 field=OxmlElement('w:fldSimple');field.set(qn('w:instr'),'PAGE');p._p.append(field)
# Remove active hyperlinks from inherited template (preserve visible runs).
for part in doc.part.package.parts:
 if hasattr(part,'_element'):
  for hyp in list(part._element.xpath('.//w:hyperlink')):
   parent=hyp.getparent();idx=parent.index(hyp)
   for child in list(hyp):parent.insert(idx,child);idx+=1
   parent.remove(hyp)
doc.core_properties.title=title;doc.core_properties.author='Joseph Igor Paredes-Ilquimiche; Alexis Gean Gerardo Vigo-Huayán';doc.core_properties.subject='Investigación experimental sobre estimación espacial de cobre y prototipo de gemelo digital';doc.core_properties.keywords='Copper Creek; Random Forest; validación por sondajes; RIAI'
doc.save(OUT/'Articulo_RIAI_Copper_Creek.docx')

# Complete Markdown delivery with the same tables, references and formulae.
md=raw
for key,(cap,heads,rows) in tables.items():
 value=cap+'\n\n|'+'|'.join(heads)+'|\n|'+'|'.join(['---']*len(heads))+'|\n'+'\n'.join('|'+'|'.join(row)+'|' for row in rows)
 md=md.replace('{{'+key+'}}',value)
latex={1:r'\bar c_j=\frac{\sum_i\ell_{ij}c_i}{\sum_i\ell_{ij}}',2:r'\hat c_{\mathrm{IDW3}}=\frac{\sum_iw_ic_i}{\sum_iw_i}',3:r'\mathrm{RMSE}=\sqrt{\frac1n\sum_i(y_i-\hat y_i)^2}',4:r'\mathrm{MAE}=\frac1n\sum_i|y_i-\hat y_i|',5:r'R^2=1-\frac{\sum_i(y_i-\hat y_i)^2}{\sum_i(y_i-\bar y)^2}'}
for n,eq in latex.items():md=md.replace('{{EQ'+str(n)+'}}','$$\n'+eq+r'\tag{'+str(n)+'}\n$$')
md=md.replace('{{REFERENCES}}','\n\n'.join(refs))
(OUT/'Articulo_RIAI_Copper_Creek.md').write_text(md,encoding='utf-8')
audit={'template_sha256':hashlib.sha256(template.read_bytes()).hexdigest(),'title_words':len(title.split()),'abstract_es_words':len(summary.split()),'abstract_en_words':len(abstract.split()),'body_words':len(body.split()),'references':len(refs),'figures':5,'tables':4,'equations':5,'docx':str(OUT/'Articulo_RIAI_Copper_Creek.docx')}
(OUT/'evidencia/maquetacion.json').write_text(json.dumps(audit,ensure_ascii=False,indent=2),encoding='utf-8');print(json.dumps(audit,ensure_ascii=False,indent=2))
